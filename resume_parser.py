"""
resume_parser.py
Extracts raw text from uploaded resume files (PDF or DOCX).

Design notes:
- We don't assume any fixed resume layout/structure (per project requirement).
- We just extract the cleanest possible raw text; structured extraction
  (skills, education, etc.) happens in a later module (skill_extractor.py),
  not here.
- Functions accept a file-like object (as returned by Streamlit's
  st.file_uploader) OR a file path, so this module works both inside the
  app and in standalone tests/notebooks.
"""

from __future__ import annotations
import io
from pathlib import Path
from typing import Union

import pdfplumber
from docx import Document

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class UnsupportedFileTypeError(Exception):
    """Raised when a file extension isn't in SUPPORTED_EXTENSIONS."""


def _get_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def extract_text_from_pdf(file_obj: Union[str, io.BytesIO]) -> str:
    """Extract text from a PDF file using pdfplumber.

    Args:
        file_obj: a file path (str) or an in-memory bytes buffer.

    Returns:
        Extracted text as a single string, with pages joined by newlines.
        Returns an empty string if no extractable text is found (e.g. the
        PDF is a scanned image with no text layer) rather than raising,
        so the caller can decide how to handle that case.
    """
    text_chunks = []
    with pdfplumber.open(file_obj) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_chunks.append(page_text)
    return "\n".join(text_chunks).strip()


def extract_text_from_docx(file_obj: Union[str, io.BytesIO]) -> str:
    """Extract text from a DOCX file using python-docx.

    Pulls text from paragraphs AND table cells, since many resumes use
    tables for layout (skills tables, two-column layouts, etc.) and a
    naive paragraph-only extraction misses that content.
    """
    document = Document(file_obj)
    text_chunks = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_chunks.append(cell.text)

    return "\n".join(text_chunks).strip()


def extract_text(filename: str, file_obj: Union[str, io.BytesIO]) -> str:
    """Dispatch to the correct extractor based on file extension.

    Args:
        filename: original filename (used only to determine extension).
        file_obj: file path or in-memory buffer to read from.

    Raises:
        UnsupportedFileTypeError: if the extension isn't .pdf or .docx.
    """
    ext = _get_extension(filename)
    if ext == ".pdf":
        return extract_text_from_pdf(file_obj)
    elif ext == ".docx":
        return extract_text_from_docx(file_obj)
    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext}'. Supported types: {SUPPORTED_EXTENSIONS}"
        )
